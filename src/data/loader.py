"""Document loader: supports PDF / DOCX / Markdown + OCR fallback"""

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from src.utils.logger import get_logger

logger = get_logger(__name__)


def _load_ocr_pages(ocr_path: Path) -> dict[int, str]:
    """Load pre-processed OCR text for scanned PDFs"""
    if not ocr_path.exists():
        return {}
    with open(ocr_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    # Convert string keys to int
    return {int(k): v for k, v in data.items()}


@dataclass
class Document:
    """Unified document data structure"""
    doc_id: str
    title: str
    content: str
    source_path: str
    doc_type: str  # pdf / docx / md
    metadata: dict = field(default_factory=dict)
    pages: list[str] = field(default_factory=list)  # text split by page
    tables: list[dict] = field(default_factory=list)  # extracted tables


class PDFLoader:
    """PDF document loader (with OCR fallback)"""

    def load(self, path: Path) -> Document:
        import fitz  # PyMuPDF

        doc = fitz.open(str(path))
        pages_text = []
        full_text = []
        tables = []

        # Pre-load OCR pages if needed
        ocr_pages = None
        processed_dir = path.parent.parent / "processed"
        if processed_dir.exists():
            doc_stem_clean = path.stem.lower().replace(' ', '').replace('-', '').replace('_', '')
            for ocr_file in processed_dir.glob("*_ocr.json"):
                ocr_stem_clean = ocr_file.stem.lower().replace('_ocr', '').replace(' ', '').replace('-', '').replace('_', '')
                if ocr_stem_clean in doc_stem_clean or doc_stem_clean in ocr_stem_clean:
                    ocr_pages = _load_ocr_pages(ocr_file)
                    logger.info(f"Loaded OCR data from {ocr_file.name} for {path.name}")
                    break

        for page_num, page in enumerate(doc):
            text = page.get_text("text")
            # If text is too short, use OCR fallback
            if len(text.strip()) < 50 and ocr_pages is not None:
                page_num_key = page_num + 1
                if page_num_key in ocr_pages:
                    text = ocr_pages[page_num_key]
                    logger.debug(f"Page {page_num+1}: using OCR text ({len(text)} chars)")
            pages_text.append(text)
            full_text.append(text)

            # Extract tables (based on PyMuPDF table detection)
            try:
                page_tables = page.find_tables()
                for table in page_tables:
                    tables.append({
                        "page": page_num + 1,
                        "data": table.extract(),
                        "bbox": table.bbox,
                    })
            except Exception:
                pass

        doc.close()
        logger.info(f"Loaded PDF: {path.name}, {len(pages_text)} pages, {len(tables)} tables")

        return Document(
            doc_id=path.stem,
            title=path.stem,
            content="\n".join(full_text),
            source_path=str(path),
            doc_type="pdf",
            pages=pages_text,
            tables=tables,
        )


class DocxLoader:
    """DOCX document loader"""

    def load(self, path: Path) -> Document:
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = "\n".join(paragraphs)

        # Extract tables
        tables = []
        for i, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                rows.append([cell.text for cell in row.cells])
            tables.append({"index": i, "data": rows})

        logger.info(f"Loaded DOCX: {path.name}, {len(paragraphs)} paragraphs, {len(tables)} tables")

        return Document(
            doc_id=path.stem,
            title=path.stem,
            content=content,
            source_path=str(path),
            doc_type="docx",
            tables=tables,
        )


class MarkdownLoader:
    """Markdown document loader"""

    def load(self, path: Path) -> Document:
        content = path.read_text(encoding="utf-8")
        logger.info(f"Loaded Markdown: {path.name}, {len(content)} chars")

        return Document(
            doc_id=path.stem,
            title=path.stem,
            content=content,
            source_path=str(path),
            doc_type="md",
        )


# File extension → loader
_LOADERS = {
    ".pdf": PDFLoader(),
    ".docx": DocxLoader(),
    ".md": MarkdownLoader(),
}


def load_document(path: str | Path) -> Document:
    """Auto-select loader based on file extension"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Document not found: {path}")
    loader = _LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    return loader.load(path)


def load_directory(dir_path: str | Path, extensions: Optional[list[str]] = None) -> list[Document]:
    """Load all documents under a directory"""
    dir_path = Path(dir_path)
    if extensions is None:
        extensions = [".pdf", ".docx", ".md"]
    docs = []
    for ext in extensions:
        for file_path in sorted(dir_path.rglob(f"*{ext}")):
            try:
                docs.append(load_document(file_path))
            except Exception as e:
                logger.warning(f"Failed to load {file_path}: {e}")
    logger.info(f"Loaded {len(docs)} documents from {dir_path}")
    return docs
